"""
Generate sheet signatures for single-sheet Excel workbooks and update Qdrant.

Usage:
    uv run python "temporary script/batch_sheet_signatures.py"

Configure batch size and dry-run behavior via the module-level constants below.
"""

from __future__ import annotations

import json
import logging
import sys
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from qdrant_client import QdrantClient

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = PROJECT_ROOT / "src"
if str(SRC_ROOT) not in sys.path:
    sys.path.insert(0, str(SRC_ROOT))

from app.config.settings import settings  # noqa: E402


# ---------------------------------------------------------------------------
# Configuration knobs (edit here instead of passing CLI args)

BATCH_SIZE = 1
SKIP = 0
DRY_RUN = True

WORKBOOK_LIST = PROJECT_ROOT / "single_sheet_workbooks.json"
WORKBOOKS_DIR = PROJECT_ROOT / "docs" / "all_sheets"


logger = logging.getLogger("sheet_signature_batch")


# ---------------------------------------------------------------------------
# Helper dataclasses and serialization utilities


@dataclass(frozen=True)
class SheetSignature:
    table_name: str
    row_count: int
    column_count: int
    columns: list[dict[str, Any]]
    sample_rows: list[dict[str, Any]]

    def to_payload(self) -> dict[str, Any]:
        return {
            "table_name": self.table_name,
            "row_count": self.row_count,
            "column_count": self.column_count,
            "columns": self.columns,
            "sample_rows": self.sample_rows,
        }


def configure_logging() -> None:
    logging.basicConfig(
        level=getattr(logging, settings.log_level.upper(), logging.INFO),
        format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
    )


def snake_case(name: str) -> str:
    base = Path(name).stem
    cleaned = (
        base.strip()
        .lower()
        .replace(" ", "_")
        .replace("-", "_")
        .replace("%", "pct")
        .replace("/", "_")
        .replace("+", "_plus_")
    )
    preserved = "".join(ch for ch in cleaned if ch.isalnum() or ch == "_")
    return "_".join(filter(None, preserved.split("_"))) or "indicator"


def point_id(indicator_name: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, indicator_name))


def load_batch() -> list[str]:
    if not WORKBOOK_LIST.exists():
        raise FileNotFoundError(WORKBOOK_LIST)
    workbook_names = json.loads(WORKBOOK_LIST.read_text())
    if not isinstance(workbook_names, list):
        raise ValueError("single_sheet_workbooks.json must contain a list")
    if not workbook_names:
        raise ValueError("single_sheet_workbooks.json is empty")
    return workbook_names[SKIP:SKIP+BATCH_SIZE]


def _row_has_values(row: Iterable[Any]) -> bool:
    return any(cell is not None and str(cell).strip() != "" for cell in row)


def _normalize_header(row: Iterable[Any]) -> list[str]:
    raw_headers = [
        (str(value).strip() if value is not None else "").strip() for value in row
    ]
    headers: list[str] = []
    seen: dict[str, int] = {}
    for idx, header in enumerate(raw_headers):
        base = header or f"column_{idx+1}"
        count = seen.get(base, 0)
        seen[base] = count + 1
        candidate = base if count == 0 else f"{base}_{count+1}"
        headers.append(candidate)
    return headers


def _infer_cell_type(value: Any) -> str:
    if isinstance(value, bool):
        return "BOOLEAN"
    if isinstance(value, (int,)):
        return "INTEGER"
    if isinstance(value, (float, Decimal)):
        return "FLOAT"
    if isinstance(value, datetime):
        return "DATETIME"
    if isinstance(value, date):
        return "DATE"
    if value is None:
        return "UNKNOWN"
    return "TEXT"


def _serialize_cell(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _collect_rows(
    rows_iter: Iterable[tuple[Any, ...]],
    headers: list[str],
) -> tuple[int, list[str], list[dict[str, Any]]]:
    row_count = 0
    column_types: list[str | None] = [None] * len(headers)
    sample_rows: list[dict[str, Any]] = []
    seen_rows: set[tuple[Any, ...]] = set()

    for row in rows_iter:
        if not _row_has_values(row):
            continue
        row_count += 1

        needs_samples = len(sample_rows) < 3
        needs_types = any(cell_type is None for cell_type in column_types)
        record: dict[str, Any] = {}

        if needs_samples or needs_types:
            for idx, header in enumerate(headers):
                cell = row[idx] if idx < len(row) else None
                value = _serialize_cell(cell)
                record[header] = value
                if column_types[idx] is None and cell is not None:
                    column_types[idx] = _infer_cell_type(cell)
        else:
            # No additional info needed; skip heavy processing.
            continue

        if needs_samples:
            signature = tuple(record.get(header) for header in headers)
            if signature not in seen_rows:
                seen_rows.add(signature)
                sample_rows.append(record)

        if needs_types:
            continue

    finalized_types = [cell_type or "TEXT" for cell_type in column_types]
    return row_count, finalized_types, sample_rows


def build_sheet_signature(indicator_name: str, workbook_path: Path) -> SheetSignature:
    if not workbook_path.exists():
        raise FileNotFoundError(workbook_path)

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet_names = wb.sheetnames
        if len(sheet_names) != 1:
            raise ValueError(
                f"Workbook {workbook_path.name} has {len(sheet_names)} sheets; expected exactly 1"
            )
        sheet = wb[sheet_names[0]]
        rows_iter = sheet.iter_rows(values_only=True)

        for header_row in rows_iter:
            if _row_has_values(header_row):
                headers = _normalize_header(header_row)
                break
        else:
            raise ValueError(f"Workbook {workbook_path.name} does not contain a header row with values")

        row_count, column_types, sample_rows = _collect_rows(rows_iter, headers)

        signature = SheetSignature(
            table_name=indicator_name,
            row_count=row_count,
            column_count=len(headers),
            columns=[
                {"name": header, "sql_type": column_types[idx]}
                for idx, header in enumerate(headers)
            ],
            sample_rows=sample_rows,
        )
        return signature
    finally:
        wb.close()


def _log_signature(signature: SheetSignature) -> None:
    payload = signature.to_payload()
    formatted = json.dumps(payload, indent=2, ensure_ascii=False)
    logger.info("Signature for %s:\n%s", signature.table_name, formatted)


def _update_qdrant_signature(client: QdrantClient, signature: SheetSignature) -> None:
    pid = point_id(signature.table_name)
    points = client.retrieve(
        collection_name=settings.qdrant_collection,
        ids=[pid],
        with_payload=True,
        with_vectors=False,
    )
    if not points:
        logger.error("No Qdrant point found for indicator %s (id=%s)", signature.table_name, pid)
        raise RuntimeError(f"Missing Qdrant point for {signature.table_name}")

    payload = points[0].payload or {}
    metadata = payload.get("metadata") or {}
    metadata["sheet_signature"] = signature.to_payload()

    client.set_payload(
        collection_name=settings.qdrant_collection,
        payload={"metadata": metadata},
        points=[pid],
    )
    logger.info(
        "Updated Qdrant payload for %s (rows=%s, cols=%s)",
        signature.table_name,
        signature.row_count,
        signature.column_count,
    )


def main() -> None:
    configure_logging()
    logger.info(
        "Starting batch | batch_size=%s | skip=%s | dry_run=%s | workbook_source=%s",
        BATCH_SIZE,
        SKIP,
        DRY_RUN,
        WORKBOOKS_DIR,
    )

    batch_files = load_batch()
    logger.info("Loaded %s workbook entries from %s", len(batch_files), WORKBOOK_LIST.name)

    client = None
    if not DRY_RUN:
        client = QdrantClient(url=settings.qdrant_url, api_key=settings.qdrant_api_key)
        logger.info("Connected to Qdrant at %s", settings.qdrant_url)

    processed = 0
    for file_name in batch_files:
        indicator_name = snake_case(file_name)
        workbook_path = WORKBOOKS_DIR / file_name
        logger.info("Processing indicator=%s | workbook=%s", indicator_name, workbook_path.name)

        signature = build_sheet_signature(indicator_name, workbook_path)
        if DRY_RUN:
            _log_signature(signature)
        else:
            if client is None:
                raise RuntimeError("Qdrant client is not initialized")
            _update_qdrant_signature(client, signature)
        processed += 1

    logger.info("Batch completed successfully | processed=%s", processed)


if __name__ == "__main__":
    main()

